"""
Router de API para el Asistente de IA en BIMEH.
Utiliza la arquitectura modular basada en IAOrchestrator, LLM Providers y SessionRepository.
"""

import os
import time
import logging
from typing import Optional, List, Dict, Any
from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel, Field

from app.database import get_db
from app.ia.orchestrator import get_ia_orchestrator
from app.ia.providers import get_llm_provider
from app.ia.repository import (
    listar_sesiones,
    crear_sesion,
    obtener_mensajes_sesion,
    eliminar_sesion,
    guardar_mensaje,
    actualizar_titulo_sesion,
    sesion_existe,
)

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/ia", tags=["IA"])

class ChatMessageRequest(BaseModel):
    message: str = Field(..., min_length=1, description="Pregunta o solicitud en lenguaje natural para la IA")
    history: Optional[List[Dict[str, Any]]] = Field(default=None, description="Historial previo de mensajes")
    active_militar: Optional[Dict[str, Any]] = Field(default=None, description="Militar en contexto activo (cedula, nombre)")
    model: Optional[str] = Field(default=None, description="Modelo de Gemini u Ollama a utilizar para esta consulta")
    mode: Optional[str] = Field(default="chat", description="Modo de operacion: 'chat' o 'report'")
    session_id: Optional[int] = Field(default=None, description="ID de la sesión de chat para persistencia")
    context: Optional[Dict[str, Any]] = Field(default=None, description="Contexto visual de pantalla y filtros activos")

class CreateSessionRequest(BaseModel):
    titulo: Optional[str] = Field(default="Nueva conversación", description="Título inicial de la sesión de chat")

class ApreciacionRequest(BaseModel):
    mes: Optional[str] = Field("TODOS", description="Nombre del mes operacional o TODOS para consolidado")
    model: Optional[str] = Field(default=None, description="Modelo de Gemini u Ollama a utilizar")

class IAConfigRequest(BaseModel):
    base_url: Optional[str] = Field(None, description="URL de Ollama o Cloudflare Tunnel (ej. https://...trycloudflare.com)")
    model: Optional[str] = Field(None, description="Modelo activo (ej. gemini-3.5-flash-lite o llama3.1)")

@router.get("/status")
def get_ia_status():
    """Reporta el estado del backend de IA activo (Gemini o Ollama) y los modelos disponibles."""
    llm, backend_name = get_llm_provider()
    return llm.get_status()

@router.get("/sesiones")
def list_chat_sessions(db = Depends(get_db)):
    """Lista las sesiones de chat persistentes ordenadas por última actualización."""
    return listar_sesiones(db)

@router.post("/sesiones")
def create_chat_session(req: Optional[CreateSessionRequest] = None, db = Depends(get_db)):
    """Crea una nueva sesión vacía de chat."""
    titulo = req.titulo if req and req.titulo else "Nueva conversación"
    return crear_sesion(db, titulo=titulo)

@router.get("/sesiones/{session_id}")
def get_chat_session(session_id: int, db = Depends(get_db)):
    """Recupera el historial de mensajes de una sesión existente."""
    session = obtener_mensajes_sesion(db, session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Sesión de chat no encontrada")
    return session

@router.delete("/sesiones/{session_id}")
def delete_chat_session(session_id: int, db = Depends(get_db)):
    """Elimina una sesión de chat y sus mensajes asociados."""
    eliminar_sesion(db, session_id)
    return {"status": "success", "message": f"Sesión {session_id} eliminada correctamente"}

@router.post("/config")
def update_ia_config(req: IAConfigRequest):
    """Actualiza la URL o modelo de IA en tiempo de ejecución (sin necesidad de reiniciar)."""
    if req.base_url:
        os.environ["OLLAMA_BASE_URL"] = req.base_url.strip().rstrip("/")
    if req.model:
        model_name = req.model.strip()
        os.environ["GEMINI_MODEL"] = model_name
        os.environ["OLLAMA_MODEL"] = model_name
        logger.info(f"[IA Config] Modelo activo actualizado a: {model_name}")
    return get_ia_status()

@router.post("/chat")
def chat_with_ia(req: ChatMessageRequest, db = Depends(get_db)):
    """
    Recibe una solicitud en lenguaje natural del usuario.
    Persiste la conversación en BD y utiliza context awareness para responder.
    """
    # 1. Resolver o crear la sesión de chat
    session_id = req.session_id
    if session_id and not sesion_existe(db, session_id):
        session_id = None

    if not session_id:
        nueva_ses = crear_sesion(db, titulo=req.message)
        session_id = nueva_ses["id"]
    else:
        # Autotitular si la sesión tenía título por defecto
        ses_data = obtener_mensajes_sesion(db, session_id)
        if ses_data and ses_data.get("titulo") in ("Nueva conversación", "Nueva consulta", ""):
            actualizar_titulo_sesion(db, session_id, req.message)

    # 2. Guardar mensaje del usuario
    guardar_mensaje(
        db=db,
        session_id=session_id,
        rol="user",
        contenido=req.message,
        metadatos={"context": req.context or {}}
    )

    # 3. Reconstruir historial si no vino del frontend
    history = req.history
    if not history:
        ses_data = obtener_mensajes_sesion(db, session_id)
        if ses_data and ses_data.get("mensajes"):
            # Excluir el último mensaje recién agregado
            prev_msgs = [m for m in ses_data["mensajes"][:-1]][-6:]
            history = [{"sender": m["rol"], "text": m["contenido"]} for m in prev_msgs]

    start_time = time.time()
    try:
        orchestrator = get_ia_orchestrator()
        resultado = orchestrator.process_query(
            user_message=req.message,
            db=db,
            history=history,
            active_militar=req.active_militar,
            model=req.model,
            mode=req.mode or "chat",
            context=req.context
        )
        elapsed = round(time.time() - start_time, 1)

        # 4. Guardar respuesta del asistente en BD
        bot_meta = {
            "elapsed_seconds": elapsed,
            "backend": resultado.get("model"),
            "model": resultado.get("model"),
            "sql": resultado.get("sql"),
            "type": resultado.get("type"),
            "detected_entities": resultado.get("detected_entities")
        }
        guardar_mensaje(
            db=db,
            session_id=session_id,
            rol="assistant",
            contenido=resultado.get("answer", ""),
            metadatos=bot_meta
        )

        return {
            "status": "success",
            "session_id": session_id,
            "elapsed_seconds": elapsed,
            "backend": resultado.get("model", "ia"),
            **resultado
        }
    except Exception as e:
        logger.error(f"[chat_with_ia] Error: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail=f"Error procesando consulta con IA: {str(e)}"
        )

@router.post("/apreciacion")
def generar_apreciacion_comandancia(req: ApreciacionRequest, db = Depends(get_db)):
    """
    Genera el informe formal de Apreciación de Situación de Personal para la comandancia.
    """
    from app.ia.skills.apreciacion_skill import ApreciacionSkill
    from app.ia.skills.base import QueryContext

    start_time = time.time()
    try:
        skill = ApreciacionSkill()
        ctx = QueryContext(user_message=f"Apreciación de personal mes {req.mes}", db=db)
        res = skill.execute(ctx)
        elapsed = round(time.time() - start_time, 1)
        if res and res.report_info:
            return {
                "status": "success",
                "elapsed_seconds": elapsed,
                **res.report_info
            }
        return {
            "status": "success",
            "elapsed_seconds": elapsed,
            "apreciacion": res.answer if res else "No se pudo generar la apreciación."
        }
    except Exception as e:
        logger.error(f"[generar_apreciacion_comandancia] Error: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail=f"Error generando apreciación: {str(e)}"
        )


# ---------------------------------------------------------------------------
# Exportación de Consultas del Asistente IA (Excel y PDF Oficial)
# ---------------------------------------------------------------------------

class ExportarTablaIARequest(BaseModel):
    titulo: str = Field(default="Reporte de Consulta de Personal", description="Título del documento")
    columns: List[str] = Field(..., description="Nombres de las columnas")
    rows: List[Dict[str, Any]] = Field(..., description="Filas con datos de la consulta")
    formato: str = Field(default="excel", description="excel o pdf")
    mes: Optional[str] = Field(default=None, description="Mes opcional de referencia")
    subtitulo: Optional[str] = Field(default=None, description="Subtítulo o descripción de la consulta")


@router.post("/exportar-tabla")
def exportar_tabla_ia(req: ExportarTablaIARequest):
    """
    Exporta el resultado de cualquier consulta del Asistente IA en Excel (.xlsx) o PDF (.pdf)
    con membrete y estándares oficiales de la unidad BIMEJ 12.
    """
    import io
    from datetime import datetime
    from fastapi.responses import StreamingResponse

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    clean_title = (req.titulo or "Reporte de Consulta").strip()

    if req.formato.lower() == "excel":
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Consulta BIMEJ12"

        # Estilos institucionales
        title_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")  # Slate 900
        title_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")

        meta_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        meta_font = Font(name="Calibri", size=9, italic=True, color="94A3B8")

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=10, bold=True, color="38BDF8")  # Cyan

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        num_cols = max(len(req.columns), 1)
        last_col_letter = get_column_letter(num_cols)

        # Fila 1: Título institucional
        ws.merge_cells(f"A1:{last_col_letter}1")
        ws["A1"] = f"BATALLÓN DE INFANTERÍA MECANIZADO N° 12 — {clean_title.upper()}"
        ws["A1"].font = title_font
        ws["A1"].fill = title_fill
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 32

        # Fila 2: Metadatos
        sub = f" · {req.subtitulo}" if req.subtitulo else ""
        ws.merge_cells(f"A2:{last_col_letter}2")
        ws["A2"] = f"BIMEJ 12 · Fecha: {now_str} · Total Registros: {len(req.rows)}{sub}"
        ws["A2"].font = meta_font
        ws["A2"].fill = meta_fill
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # Fila 3: Espacio
        ws.row_dimensions[3].height = 10

        # Fila 4: Cabeceras de columnas
        ws.row_dimensions[4].height = 24
        for col_idx, col_name in enumerate(req.columns, start=1):
            cell = ws.cell(row=4, column=col_idx)
            cell.value = str(col_name).upper().replace("_", " ")
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Filas 5 en adelante: Datos
        data_font = Font(name="Calibri", size=10)
        mono_font = Font(name="Consolas", size=9)
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        for r_idx, row in enumerate(req.rows, start=5):
            ws.row_dimensions[r_idx].height = 20
            is_even = (r_idx % 2 == 0)
            for c_idx, col_name in enumerate(req.columns, start=1):
                val = row.get(col_name, "")
                if val is None:
                    val = "-"
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = str(val) if not isinstance(val, (int, float)) else val
                cell.border = thin_border
                if is_even:
                    cell.fill = zebra_fill

                col_lower = str(col_name).lower()
                if col_lower in ("cedula", "fecha_reporte", "desde", "hasta", "fecha"):
                    cell.font = mono_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_lower in ("dias", "total_dias", "id", "cant"):
                    cell.font = data_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = data_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Ajuste de ancho de columnas automático
        for col_idx, col_name in enumerate(req.columns, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = len(str(col_name))
            for row in req.rows[:100]:
                val = row.get(col_name)
                if val:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)

        filename = f"Reporte_Consulta_BIMEJ12_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return StreamingResponse(
            out,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )

    elif req.formato.lower() == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

        out = io.BytesIO()
        page_layout = landscape(letter) if len(req.columns) > 4 else letter

        doc = SimpleDocTemplate(
            out,
            pagesize=page_layout,
            leftMargin=28,
            rightMargin=28,
            topMargin=30,
            bottomMargin=30
        )

        styles = getSampleStyleSheet()
        uid = id(out)

        title_style = ParagraphStyle(
            f'TitleStyle_{uid}',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=12,
            textColor=colors.HexColor('#0F172A'),
            alignment=1,
            spaceAfter=3
        )

        subtitle_style = ParagraphStyle(
            f'SubtitleStyle_{uid}',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            textColor=colors.HexColor('#475569'),
            alignment=1,
            spaceAfter=10
        )

        cell_style = ParagraphStyle(
            f'CellStyle_{uid}',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            textColor=colors.HexColor('#1E293B'),
            leading=10
        )

        header_cell_style = ParagraphStyle(
            f'HeaderCellStyle_{uid}',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            textColor=colors.white,
            alignment=1,
            leading=10
        )

        elements = []
        elements.append(Paragraph("<b>EJÉRCITO NACIONAL DE COLOMBIA</b>", subtitle_style))
        elements.append(Paragraph("BATALLÓN DE INFANTERÍA MECANIZADO N° 12 'GR. RAMÓN ARTURO RINCÓN QUIÑÓNEZ'", title_style))
        sub = f" — {req.subtitulo}" if req.subtitulo else ""
        elements.append(Paragraph(f"<b>{clean_title.upper()}</b>{sub} · Fecha: {now_str} · Total registros: {len(req.rows)}", subtitle_style))
        elements.append(Spacer(1, 6))

        # Construir matriz de tabla
        table_data = []

        headers = [Paragraph(f"<b>{str(c).upper().replace('_', ' ')}</b>", header_cell_style) for c in req.columns]
        headers.insert(0, Paragraph("<b>#</b>", header_cell_style))
        table_data.append(headers)

        for idx, row in enumerate(req.rows, start=1):
            row_cells = [Paragraph(str(idx), cell_style)]
            for col in req.columns:
                val = row.get(col, "")
                if val is None or val == "":
                    val = "-"
                row_cells.append(Paragraph(str(val), cell_style))
            table_data.append(row_cells)

        t = Table(table_data, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#94A3B8')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('TOPPADDING', (0, 0), (-1, -1), 3.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
        ]))
        elements.append(t)

        doc.build(elements)
        out.seek(0)

        filename = f"Reporte_Consulta_BIMEJ12_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        return StreamingResponse(
            out,
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )

    else:
        raise HTTPException(status_code=400, detail="Formato no soportado. Use 'excel' o 'pdf'.")

