from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from typing import List, Optional, Dict, Any
from datetime import datetime, date, timedelta
import io
from collections import defaultdict
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import letter, landscape

from app.database import get_db, get_month_dates

router = APIRouter(prefix="/api/novedades", tags=["GestionNovedades"])


@router.get("/catalogo")
def get_catalogo_novedades(db = Depends(get_db)):
    """
    Retorna la lista completa de sub_novedades con el conteo acumulado de registros en la base de datos.
    """
    cursor = db.cursor()
    cursor.execute("""
        SELECT sn.id, sn.nombre, COUNT(rp.id) AS total_registros, COUNT(DISTINCT rp.id_personal) AS total_personal
        FROM SUB_NOVEDADES sn
        LEFT JOIN REGISTRO_PERSONAL rp ON sn.id = rp.id_sub_novedad
        GROUP BY sn.id, sn.nombre
        ORDER BY sn.id ASC;
    """)
    rows = cursor.fetchall()
    return [
        {
            "id": row[0],
            "nombre": row[1],
            "total_registros": row[2],
            "total_personal": row[3]
        }
        for row in rows
    ]


def _clean_param(v):
    from fastapi.params import Query as QueryParam
    if isinstance(v, QueryParam):
        return v.default
    return v

def _build_novedades_where_clause(
    id_sub_novedad: Optional[int] = None,
    mes: Optional[str] = None,
    fecha_inicio: Optional[str] = None,
    fecha_fin: Optional[str] = None,
    q: Optional[str] = None,
    estado: Optional[str] = None,
):
    id_sub_novedad = _clean_param(id_sub_novedad)
    mes = _clean_param(mes)
    fecha_inicio = _clean_param(fecha_inicio)
    fecha_fin = _clean_param(fecha_fin)
    q = _clean_param(q)
    estado = _clean_param(estado)

    # 1. Asegurar consistencia de orden en fechas (evitar rangos invertidos como inicio > fin)
    if fecha_inicio and fecha_fin and str(fecha_inicio).strip() and str(fecha_fin).strip():
        fecha_inicio_str = str(fecha_inicio).strip()
        fecha_fin_str = str(fecha_fin).strip()
        if fecha_inicio_str > fecha_fin_str:
            fecha_inicio, fecha_fin = fecha_fin_str, fecha_inicio_str
        else:
            fecha_inicio = fecha_inicio_str
            fecha_fin = fecha_fin_str
    elif fecha_inicio:
        fecha_inicio = str(fecha_inicio).strip()
    elif fecha_fin:
        fecha_fin = str(fecha_fin).strip()

    where_clauses = []
    params = []

    if id_sub_novedad is not None and int(id_sub_novedad) > 0:
        where_clauses.append("rp.id_sub_novedad = %s")
        params.append(int(id_sub_novedad))

    if mes and str(mes).upper() != "TODOS":
        dates = get_month_dates(mes)
        if dates:
            # Si se pasó fecha_inicio o fecha_fin junto a un mes específico, 
            # filtrar los dates del mes que se encuentren dentro del rango solicitado
            if fecha_inicio:
                dates = [d for d in dates if d >= fecha_inicio]
            if fecha_fin:
                dates = [d for d in dates if d <= fecha_fin]
            
            if dates:
                placeholders = ",".join("%s" for _ in dates)
                where_clauses.append(f"r.fecha IN ({placeholders})")
                params.extend(dates)
            else:
                where_clauses.append("1=0")
        else:
            where_clauses.append("1=0")
    else:
        if fecha_inicio:
            where_clauses.append("r.fecha >= %s")
            params.append(fecha_inicio)

        if fecha_fin:
            where_clauses.append("r.fecha <= %s")
            params.append(fecha_fin)

    if q and q.strip():
        search_pattern = f"%{q.strip().upper()}%"
        where_clauses.append("(CAST(p.cedula AS TEXT) LIKE %s OR UPPER(p.nombre) LIKE %s)")
        params.extend([search_pattern, search_pattern])

    if estado and estado.upper() != "TODOS":
        if estado.upper() == "ACTIVO":
            where_clauses.append("p.fecha_retiro IS NULL")
        elif estado.upper() == "RETIRADO":
            where_clauses.append("p.fecha_retiro IS NOT NULL")

    where_sql = ("WHERE " + " AND ".join(where_clauses)) if where_clauses else ""
    return where_sql, params


@router.get("/consulta")
def consultar_novedades(
    id_sub_novedad: Optional[int] = Query(None, description="ID de la subnovedad a consultar"),
    mes: Optional[str] = Query(None, description="Nombre del mes (ej. ENERO, FEBRERO) o TODOS"),
    fecha_inicio: Optional[str] = Query(None, description="Fecha mínima YYYY-MM-DD"),
    fecha_fin: Optional[str] = Query(None, description="Fecha máxima YYYY-MM-DD"),
    q: Optional[str] = Query(None, description="Búsqueda por cédula o nombre"),
    estado: Optional[str] = Query("TODOS", description="TODOS, ACTIVO o RETIRADO"),
    page: int = Query(1, ge=1),
    limit: int = Query(50, ge=1, le=500),
    db = Depends(get_db)
):
    cursor = db.cursor()
    where_sql, params = _build_novedades_where_clause(
        id_sub_novedad=id_sub_novedad,
        mes=mes,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
        q=q,
        estado=estado
    )

    # 1. Calcular KPIs globales con los filtros aplicados
    kpi_query = f"""
        SELECT 
            COUNT(rp.id) AS total_registros,
            COUNT(DISTINCT rp.id_personal) AS personal_unico,
            COUNT(DISTINCT CASE WHEN p.fecha_retiro IS NULL THEN p.id ELSE NULL END) AS activos_unicos,
            COUNT(DISTINCT CASE WHEN p.fecha_retiro IS NOT NULL THEN p.id ELSE NULL END) AS retirados_unicos,
            MIN(r.fecha) AS primera_fecha,
            MAX(r.fecha) AS ultima_fecha
        FROM REGISTRO_PERSONAL rp
        JOIN PERSONAL p ON rp.id_personal = p.id
        JOIN SUB_NOVEDADES sn ON rp.id_sub_novedad = sn.id
        JOIN REPORTES r ON rp.id_reporte = r.id
        {where_sql};
    """
    cursor.execute(kpi_query, params)
    kpi_row = cursor.fetchone()

    total_registros = kpi_row[0] or 0
    personal_unico = kpi_row[1] or 0
    activos_unicos = kpi_row[2] or 0
    retirados_unicos = kpi_row[3] or 0
    primera_fecha = kpi_row[4]
    ultima_fecha = kpi_row[5]

    # 2. Paginación de resultados
    offset = (page - 1) * limit
    data_query = f"""
        SELECT 
            rp.id,
            p.cedula,
            p.nombre,
            CASE WHEN p.fecha_retiro IS NULL THEN 'ACTIVO' ELSE 'RETIRADO' END AS estado,
            p.fecha_retiro,
            sn.id AS id_sub_novedad,
            sn.nombre AS sub_novedad,
            r.fecha AS fecha_reporte,
            COALESCE(rp.fecha_inicio, '') AS fecha_inicio,
            COALESCE(rp.fecha_final, '') AS fecha_final,
            COALESCE(rp.descripcion, '') AS descripcion
        FROM REGISTRO_PERSONAL rp
        JOIN PERSONAL p ON rp.id_personal = p.id
        JOIN SUB_NOVEDADES sn ON rp.id_sub_novedad = sn.id
        JOIN REPORTES r ON rp.id_reporte = r.id
        {where_sql}
        ORDER BY r.fecha DESC, p.nombre ASC
        LIMIT %s OFFSET %s;
    """
    cursor.execute(data_query, params + [limit, offset])
    rows = cursor.fetchall()

    registros = [
        {
            "id": row[0],
            "cedula": row[1],
            "nombre": row[2],
            "estado": row[3],
            "fecha_retiro": row[4],
            "id_sub_novedad": row[5],
            "sub_novedad": row[6],
            "fecha_reporte": row[7],
            "fecha_inicio": row[8],
            "fecha_final": row[9],
            "descripcion": row[10]
        }
        for row in rows
    ]

    total_pages = (total_registros + limit - 1) // limit if limit > 0 else 1

    return {
        "kpis": {
            "total_registros": total_registros,
            "personal_unico": personal_unico,
            "activos_unicos": activos_unicos,
            "retirados_unicos": retirados_unicos,
            "primera_fecha": primera_fecha,
            "ultima_fecha": ultima_fecha
        },
        "registros": registros,
        "total": total_registros,
        "page": page,
        "limit": limit,
        "total_pages": total_pages
    }


def format_agil_date_ranges(records: List[tuple], highlight_html: bool = False) -> str:
    """
    records: lista de tuplas (date_str, subnovedad_str, [desc_str])
    Comprime días calendario continuos de una misma novedad en rangos legibles:
    - Dentro del mismo mes: '01-10 (VACACIONES)' o '15 (PERMISO)'
    - A través de meses: '28/05-04/06 (INCAPACIDAD)'
    Si highlight_html=True: envuelve las fechas en <font color="#DC2626"><b>...</b></font> para reportes PDF.
    """
    if not records:
        return "-"

    parsed = []
    for r in records:
        d_str = str(r[0])
        nov = str(r[1])
        desc = str(r[2]) if len(r) > 2 and r[2] else ""
        try:
            dt = datetime.strptime(d_str[:10], "%Y-%m-%d").date()
            parsed.append((dt, nov, desc))
        except Exception:
            continue

    if not parsed:
        return "-"

    parsed.sort(key=lambda x: x[0])
    single_month = len(set((p[0].year, p[0].month) for p in parsed)) == 1

    ranges = []
    curr_start, curr_nov, _ = parsed[0]
    curr_end = curr_start

    def make_label(start_d: date, end_d: date, nov_name: str) -> str:
        if single_month:
            day_str = f"{start_d.day:02d}" if start_d == end_d else f"{start_d.day:02d}-{end_d.day:02d}"
        else:
            if start_d == end_d:
                day_str = f"{start_d.day:02d}/{start_d.month:02d}"
            elif start_d.month == end_d.month:
                day_str = f"{start_d.day:02d}-{end_d.day:02d}/{start_d.month:02d}"
            else:
                day_str = f"{start_d.day:02d}/{start_d.month:02d}-{end_d.day:02d}/{end_d.month:02d}"

        if highlight_html:
            return f'<font color="#DC2626"><b>{day_str}</b></font> ({nov_name})'
        else:
            return f"{day_str} ({nov_name})"

    for p in parsed[1:]:
        d, nov, _ = p
        if d == curr_end + timedelta(days=1) and nov == curr_nov:
            curr_end = d
        else:
            ranges.append(make_label(curr_start, curr_end, curr_nov))
            curr_start = d
            curr_end = d
            curr_nov = nov

    ranges.append(make_label(curr_start, curr_end, curr_nov))
    sep = "<br/>" if highlight_html else "\n"
    return sep.join(ranges)


@router.get("/builder/preview")
def builder_preview(
    modo: str = Query("agil", description="detallado | agil"),
    id_sub_novedad: Optional[int] = Query(None),
    mes: Optional[str] = Query(None),
    fecha_inicio: Optional[str] = Query(None),
    fecha_fin: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    estado: Optional[str] = Query("TODOS"),
    min_dias: int = Query(0),
    columnas: Optional[str] = Query(None),
    orden: str = Query("nombre_asc"),
    db = Depends(get_db)
):
    modo = str(_clean_param(modo) or "agil")
    id_sub_novedad = _clean_param(id_sub_novedad)
    mes = _clean_param(mes)
    fecha_inicio = _clean_param(fecha_inicio)
    fecha_fin = _clean_param(fecha_fin)
    q = _clean_param(q)
    estado = _clean_param(estado)
    min_dias = int(_clean_param(min_dias) or 0)
    orden = str(_clean_param(orden) or "nombre_asc")
    columnas = _clean_param(columnas)

    cursor = db.cursor()
    where_sql, params = _build_novedades_where_clause(
        id_sub_novedad=id_sub_novedad,
        mes=mes,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
        q=q,
        estado=estado
    )

    query = f"""
        SELECT 
            p.cedula,
            p.nombre,
            CASE WHEN p.fecha_retiro IS NULL THEN 'ACTIVO' ELSE 'RETIRADO' END AS estado,
            sn.nombre AS sub_novedad,
            r.fecha AS fecha_reporte,
            COALESCE(rp.fecha_inicio, '') AS fecha_inicio,
            COALESCE(rp.fecha_final, '') AS fecha_final,
            COALESCE(rp.descripcion, '') AS descripcion
        FROM REGISTRO_PERSONAL rp
        JOIN PERSONAL p ON rp.id_personal = p.id
        JOIN SUB_NOVEDADES sn ON rp.id_sub_novedad = sn.id
        JOIN REPORTES r ON rp.id_reporte = r.id
        {where_sql}
        ORDER BY r.fecha ASC, p.nombre ASC;
    """
    cursor.execute(query, params)
    rows = cursor.fetchall()

    if modo == "agil":
        # Agrupar por integrante (cero filas duplicadas)
        person_dict = defaultdict(list)
        for r in rows:
            person_key = (r[0], r[1], r[2]) # (cedula, nombre, estado)
            person_dict[person_key].append(r)

        items_list = []
        for (cedula, nombre, p_estado), p_rows in person_dict.items():
            dias_acumulados = len(p_rows)
            if min_dias > 0 and dias_acumulados < min_dias:
                continue

            p_rows.sort(key=lambda x: x[4])
            recs_for_ranges = [(x[4], x[3], x[7]) for x in p_rows]
            rango_fechas = format_agil_date_ranges(recs_for_ranges, highlight_html=False)

            subnovs = sorted(list(set(x[3] for x in p_rows)))
            subnov_str = ", ".join(subnovs)

            unique_descs = [x[7].strip() for x in p_rows if x[7] and str(x[7]).strip()]
            seen_descs = []
            for d in unique_descs:
                if d not in seen_descs:
                    seen_descs.append(d)
            desc_str = "; ".join(seen_descs[:3]) if seen_descs else "-"

            items_list.append({
                "cedula": cedula,
                "nombre": nombre,
                "estado": p_estado,
                "sub_novedad": subnov_str,
                "rango_fechas": rango_fechas,
                "dias_acumulados": dias_acumulados,
                "descripcion": desc_str,
                "primera_fecha": p_rows[0][4],
                "ultima_fecha": p_rows[-1][4]
            })

        # Ordenar
        if orden == "nombre_asc":
            items_list.sort(key=lambda x: x["nombre"])
        elif orden == "dias_desc":
            items_list.sort(key=lambda x: x["dias_acumulados"], reverse=True)
        elif orden == "fecha_desc":
            items_list.sort(key=lambda x: x["ultima_fecha"], reverse=True)
        elif orden == "fecha_asc":
            items_list.sort(key=lambda x: x["primera_fecha"])

        total_filas = len(items_list)
        total_personal = len(items_list)
        preview_filas = items_list[:5]

    else:
        # Modo detallado (1 fila por cada reporte diario)
        if min_dias > 0:
            person_counts = defaultdict(int)
            for r in rows:
                person_counts[r[0]] += 1
            rows = [r for r in rows if person_counts[r[0]] >= min_dias]

        if orden == "nombre_asc":
            rows.sort(key=lambda x: (x[1], x[4]))
        elif orden == "dias_desc":
            person_counts = defaultdict(int)
            for r in rows:
                person_counts[r[0]] += 1
            rows.sort(key=lambda x: (person_counts[x[0]], x[1], x[4]), reverse=True)
        elif orden == "fecha_desc":
            rows.sort(key=lambda x: (x[4], x[1]), reverse=True)
        elif orden == "fecha_asc":
            rows.sort(key=lambda x: (x[4], x[1]))

        total_filas = len(rows)
        total_personal = len(set(r[0] for r in rows))
        preview_filas = [
            {
                "cedula": r[0],
                "nombre": r[1],
                "estado": r[2],
                "sub_novedad": r[3],
                "fecha_reporte": r[4],
                "fecha_inicio": r[5] if r[5] else "-",
                "fecha_final": r[6] if r[6] else "-",
                "descripcion": r[7] if r[7] else "-"
            }
            for r in rows[:5]
        ]

    return {
        "modo": modo,
        "total_filas": total_filas,
        "total_personal": total_personal,
        "filas": preview_filas
    }


@router.get("/exportar/excel")
def exportar_novedades_excel(
    modo: str = Query("detallado", description="detallado | agil"),
    id_sub_novedad: Optional[int] = Query(None),
    mes: Optional[str] = Query(None),
    fecha_inicio: Optional[str] = Query(None),
    fecha_fin: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    estado: Optional[str] = Query("TODOS"),
    min_dias: int = Query(0),
    columnas: Optional[str] = Query(None),
    orden: str = Query("nombre_asc"),
    db = Depends(get_db)
):
    modo = str(_clean_param(modo) or "detallado")
    id_sub_novedad = _clean_param(id_sub_novedad)
    mes = _clean_param(mes)
    fecha_inicio = _clean_param(fecha_inicio)
    fecha_fin = _clean_param(fecha_fin)
    q = _clean_param(q)
    estado = _clean_param(estado)
    min_dias = int(_clean_param(min_dias) or 0)
    orden = str(_clean_param(orden) or "nombre_asc")
    columnas = _clean_param(columnas)

    cursor = db.cursor()
    where_sql, params = _build_novedades_where_clause(
        id_sub_novedad=id_sub_novedad,
        mes=mes,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
        q=q,
        estado=estado
    )

    subnovedad_nombre = "TODAS LAS NOVEDADES"
    if id_sub_novedad:
        cursor.execute("SELECT nombre FROM SUB_NOVEDADES WHERE id = %s;", (id_sub_novedad,))
        s_row = cursor.fetchone()
        if s_row:
            subnovedad_nombre = s_row[0]

    query = f"""
        SELECT 
            p.cedula,
            p.nombre,
            CASE WHEN p.fecha_retiro IS NULL THEN 'ACTIVO' ELSE 'RETIRADO' END AS estado,
            sn.nombre AS sub_novedad,
            r.fecha AS fecha_reporte,
            COALESCE(rp.fecha_inicio, '') AS fecha_inicio,
            COALESCE(rp.fecha_final, '') AS fecha_final,
            COALESCE(rp.descripcion, '') AS descripcion
        FROM REGISTRO_PERSONAL rp
        JOIN PERSONAL p ON rp.id_personal = p.id
        JOIN SUB_NOVEDADES sn ON rp.id_sub_novedad = sn.id
        JOIN REPORTES r ON rp.id_reporte = r.id
        {where_sql}
        ORDER BY r.fecha ASC, p.nombre ASC
        LIMIT 25000;
    """
    cursor.execute(query, params)
    rows = cursor.fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.views.sheetView[0].showGridLines = True

    # Estilos institucionales
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    th_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    th_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    if modo == "agil":
        ws.title = "Reporte Ágil de Novedades"
        
        # Agrupar registros por integrante
        person_dict = defaultdict(list)
        for r in rows:
            person_key = (r[0], r[1], r[2])
            person_dict[person_key].append(r)

        items_list = []
        for (cedula, nombre, p_estado), p_rows in person_dict.items():
            dias_acumulados = len(p_rows)
            if min_dias > 0 and dias_acumulados < min_dias:
                continue

            p_rows.sort(key=lambda x: x[4])
            recs_for_ranges = [(x[4], x[3], x[7]) for x in p_rows]
            rango_fechas = format_agil_date_ranges(recs_for_ranges, highlight_html=False)

            subnovs = sorted(list(set(x[3] for x in p_rows)))
            subnov_str = ", ".join(subnovs)

            unique_descs = [f"• {x[4]} ({x[3]}): {x[7].strip()}" for x in p_rows if x[7] and str(x[7]).strip()]
            desc_str = "; ".join([x[7].strip() for x in p_rows if x[7] and str(x[7]).strip()][:3]) or "-"

            items_list.append({
                "cedula": cedula,
                "nombre": nombre,
                "estado": p_estado,
                "sub_novedad": subnov_str,
                "rango_fechas": rango_fechas,
                "dias_acumulados": dias_acumulados,
                "descripcion": desc_str,
                "raw_recs": p_rows,
                "observaciones_list": unique_descs,
                "primera_fecha": p_rows[0][4],
                "ultima_fecha": p_rows[-1][4]
            })

        if orden == "nombre_asc":
            items_list.sort(key=lambda x: x["nombre"])
        elif orden == "dias_desc":
            items_list.sort(key=lambda x: x["dias_acumulados"], reverse=True)
        elif orden == "fecha_desc":
            items_list.sort(key=lambda x: x["ultima_fecha"], reverse=True)
        elif orden == "fecha_asc":
            items_list.sort(key=lambda x: x["primera_fecha"])

        # 1. Encabezado institucional
        ws.merge_cells("A1:G1")
        ws["A1"] = "BIMEJ 12 - BATALLÓN DE INGENIEROS DE MOVILIDAD Y CONTRAMOVILIDAD N° 12"
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="0F172A")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25

        ws.merge_cells("A2:G2")
        ws["A2"] = "REPORTE ÁGIL DE NOVEDADES - RESUMEN PARA COMANDANCIA (SIN DUPLICADOS)"
        ws["A2"].font = Font(name="Calibri", size=11, bold=True, color="0284C7")
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # 2. Bloque de metadatos del reporte
        periodo_label = mes if mes and mes != "TODOS" else "AÑO COMPLETO"
        if fecha_inicio or fecha_fin:
            periodo_label = f"{fecha_inicio or 'Inicio'} a {fecha_fin or 'Actual'}"

        metadata = [
            ("Novedad Auditada:", subnovedad_nombre, "Total Efectivos:", len(items_list)),
            ("Periodo:", periodo_label, "Días Acumulados:", sum(it["dias_acumulados"] for it in items_list)),
            ("Estado Personal:", estado, "Fecha Emisión:", datetime.now().strftime("%d/%m/%Y %H:%M"))
        ]

        curr_row = 4
        for f1_label, f1_val, f2_label, f2_val in metadata:
            ws.cell(row=curr_row, column=1, value=f1_label).font = Font(name="Calibri", size=10, bold=True, color="475569")
            ws.cell(row=curr_row, column=2, value=str(f1_val)).font = Font(name="Calibri", size=10, bold=True, color="0F172A")
            ws.cell(row=curr_row, column=4, value=f2_label).font = Font(name="Calibri", size=10, bold=True, color="475569")
            ws.cell(row=curr_row, column=5, value=str(f2_val)).font = Font(name="Calibri", size=10, bold=True, color="0F172A")
            curr_row += 1

        curr_row += 1

        # 3. Encabezados de tabla Ágil
        headers = [
            "CÉDULA",
            "APELLIDOS Y NOMBRES",
            "ESTADO",
            "NOVEDAD(ES)",
            "RANGOS CONDENSADOS",
            "DÍAS",
            "OBSERVACIONES / MOTIVO"
        ]
        col_widths = [14, 34, 12, 22, 28, 10, 42]

        for col_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h_text)
            cell.font = th_font
            cell.fill = th_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[curr_row].height = 24
        curr_row += 1

        # 4. Filas de datos
        for idx, item in enumerate(items_list):
            row_fill = zebra_fill if idx % 2 == 1 else white_fill
            c1 = ws.cell(row=curr_row, column=1, value=item["cedula"])
            c2 = ws.cell(row=curr_row, column=2, value=item["nombre"])
            c3 = ws.cell(row=curr_row, column=3, value=item["estado"])
            c4 = ws.cell(row=curr_row, column=4, value=item["sub_novedad"])
            c5 = ws.cell(row=curr_row, column=5, value=item["rango_fechas"])
            c6 = ws.cell(row=curr_row, column=6, value=item["dias_acumulados"])
            c7 = ws.cell(row=curr_row, column=7, value=item["descripcion"])

            # Comentario emergente en la celda de rangos si hay descripciones
            if item["observaciones_list"]:
                comm_text = "HISTORIAL Y OBSERVACIONES DETALLADAS:\n" + "\n".join(item["observaciones_list"][:15])
                c5.comment = Comment(comm_text, "BIMEJ12")

            cells = [c1, c2, c3, c4, c5, c6, c7]
            for c_idx, cell in enumerate(cells, 1):
                cell.font = Font(name="Calibri", size=9)
                cell.fill = row_fill
                cell.border = thin_border
                if c_idx in (1, 3, 6):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 5:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            # Altura dinámica según líneas en rangos condensados
            num_lines = item["rango_fechas"].count("\n") + 1
            ws.row_dimensions[curr_row].height = max(20, num_lines * 17)
            curr_row += 1

        for i, w in enumerate(col_widths, 1):
            col_letter = get_column_letter(i)
            ws.column_dimensions[col_letter].width = w

        filename = f"Reporte_Agil_Novedades_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    else:
        # Modo detallado (1 fila por cada jornada)
        ws.title = "Reporte Detallado de Novedades"

        if min_dias > 0:
            person_counts = defaultdict(int)
            for r in rows:
                person_counts[r[0]] += 1
            rows = [r for r in rows if person_counts[r[0]] >= min_dias]

        if orden == "nombre_asc":
            rows.sort(key=lambda x: (x[1], x[4]))
        elif orden == "dias_desc":
            person_counts = defaultdict(int)
            for r in rows:
                person_counts[r[0]] += 1
            rows.sort(key=lambda x: (person_counts[x[0]], x[1], x[4]), reverse=True)
        elif orden == "fecha_desc":
            rows.sort(key=lambda x: (x[4], x[1]), reverse=True)
        elif orden == "fecha_asc":
            rows.sort(key=lambda x: (x[4], x[1]))

        ws.merge_cells("A1:H1")
        ws["A1"] = "BIMEJ 12 - BATALLÓN DE INGENIEROS DE MOVILIDAD Y CONTRAMOVILIDAD N° 12"
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="0F172A")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25

        ws.merge_cells("A2:H2")
        ws["A2"] = "REPORTE DETALLADO DE NOVEDADES DE PERSONAL (AUDITORÍA COMPLETA)"
        ws["A2"].font = Font(name="Calibri", size=11, bold=True, color="0284C7")
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        periodo_label = mes if mes and mes != "TODOS" else "AÑO COMPLETO"
        if fecha_inicio or fecha_fin:
            periodo_label = f"{fecha_inicio or 'Inicio'} a {fecha_fin or 'Actual'}"

        metadata = [
            ("Novedad Auditada:", subnovedad_nombre, "Total Registros:", len(rows)),
            ("Periodo:", periodo_label, "Personal Único:", len(set(r[0] for r in rows))),
            ("Estado Personal:", estado, "Fecha Emisión:", datetime.now().strftime("%d/%m/%Y %H:%M"))
        ]

        curr_row = 4
        for f1_label, f1_val, f2_label, f2_val in metadata:
            ws.cell(row=curr_row, column=1, value=f1_label).font = Font(name="Calibri", size=10, bold=True, color="475569")
            ws.cell(row=curr_row, column=2, value=str(f1_val)).font = Font(name="Calibri", size=10, bold=True, color="0F172A")
            ws.cell(row=curr_row, column=4, value=f2_label).font = Font(name="Calibri", size=10, bold=True, color="475569")
            ws.cell(row=curr_row, column=5, value=str(f2_val)).font = Font(name="Calibri", size=10, bold=True, color="0F172A")
            curr_row += 1

        curr_row += 1

        headers = [
            "CÉDULA",
            "APELLIDOS Y NOMBRES",
            "ESTADO",
            "NOVEDAD ASIGNADA",
            "FECHA REPORTE",
            "FECHA INICIO",
            "FECHA FINAL",
            "DESCRIPCIÓN / OBSERVACIÓN"
        ]
        col_widths = [14, 34, 12, 25, 15, 14, 14, 40]

        for col_idx, header_text in enumerate(headers, 1):
            cell = ws.cell(row=curr_row, column=col_idx, value=header_text)
            cell.font = th_font
            cell.fill = th_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[curr_row].height = 24
        curr_row += 1

        for idx, r in enumerate(rows):
            row_fill = zebra_fill if idx % 2 == 1 else white_fill
            r_cells = [
                ws.cell(row=curr_row, column=1, value=r[0]),
                ws.cell(row=curr_row, column=2, value=r[1]),
                ws.cell(row=curr_row, column=3, value=r[2]),
                ws.cell(row=curr_row, column=4, value=r[3]),
                ws.cell(row=curr_row, column=5, value=r[4]),
                ws.cell(row=curr_row, column=6, value=r[5] if r[5] else "-"),
                ws.cell(row=curr_row, column=7, value=r[6] if r[6] else "-"),
                ws.cell(row=curr_row, column=8, value=r[7] if r[7] else "-")
            ]

            for c_idx, cell in enumerate(r_cells, 1):
                cell.font = Font(name="Calibri", size=9)
                cell.fill = row_fill
                cell.border = thin_border
                if c_idx in (1, 3, 5, 6, 7):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws.row_dimensions[curr_row].height = 18
            curr_row += 1

        for i, w in enumerate(col_widths, 1):
            col_letter = get_column_letter(i)
            ws.column_dimensions[col_letter].width = w

        filename = f"Reporte_Detallado_Novedades_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/exportar/pdf")
def exportar_novedades_pdf(
    modo: str = Query("detallado", description="detallado | agil"),
    id_sub_novedad: Optional[int] = Query(None),
    mes: Optional[str] = Query(None),
    fecha_inicio: Optional[str] = Query(None),
    fecha_fin: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    estado: Optional[str] = Query("TODOS"),
    min_dias: int = Query(0),
    columnas: Optional[str] = Query(None),
    orden: str = Query("nombre_asc"),
    db = Depends(get_db)
):
    modo = str(_clean_param(modo) or "detallado")
    id_sub_novedad = _clean_param(id_sub_novedad)
    mes = _clean_param(mes)
    fecha_inicio = _clean_param(fecha_inicio)
    fecha_fin = _clean_param(fecha_fin)
    q = _clean_param(q)
    estado = _clean_param(estado)
    min_dias = int(_clean_param(min_dias) or 0)
    orden = str(_clean_param(orden) or "nombre_asc")
    columnas = _clean_param(columnas)

    cursor = db.cursor()
    where_sql, params = _build_novedades_where_clause(
        id_sub_novedad=id_sub_novedad,
        mes=mes,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
        q=q,
        estado=estado
    )

    subnovedad_nombre = "TODAS LAS NOVEDADES"
    if id_sub_novedad:
        cursor.execute("SELECT nombre FROM SUB_NOVEDADES WHERE id = %s;", (id_sub_novedad,))
        s_row = cursor.fetchone()
        if s_row:
            subnovedad_nombre = s_row[0]

    query = f"""
        SELECT 
            p.cedula,
            p.nombre,
            CASE WHEN p.fecha_retiro IS NULL THEN 'ACTIVO' ELSE 'RETIRADO' END AS estado,
            sn.nombre AS sub_novedad,
            r.fecha AS fecha_reporte,
            COALESCE(rp.fecha_inicio, '') AS fecha_inicio,
            COALESCE(rp.fecha_final, '') AS fecha_final,
            COALESCE(rp.descripcion, '') AS descripcion
        FROM REGISTRO_PERSONAL rp
        JOIN PERSONAL p ON rp.id_personal = p.id
        JOIN SUB_NOVEDADES sn ON rp.id_sub_novedad = sn.id
        JOIN REPORTES r ON rp.id_reporte = r.id
        {where_sql}
        ORDER BY r.fecha ASC, p.nombre ASC
        LIMIT 10000;
    """
    cursor.execute(query, params)
    rows = cursor.fetchall()

    pdf_buffer = io.BytesIO()
    uid = id(pdf_buffer)
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=landscape(letter),
        leftMargin=25,
        rightMargin=25,
        topMargin=25,
        bottomMargin=25
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        f'NovedadTitle_{uid}',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=14,
        textColor=colors.HexColor('#0F172A'),
        alignment=1,
        spaceAfter=3
    )
    subtitle_style = ParagraphStyle(
        f'NovedadSubtitle_{uid}',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        textColor=colors.HexColor('#0284C7'),
        alignment=1,
        spaceAfter=10
    )
    meta_key_style = ParagraphStyle(
        f'MetaKey_{uid}',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=colors.HexColor('#475569')
    )
    meta_val_style = ParagraphStyle(
        f'MetaVal_{uid}',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        textColor=colors.HexColor('#0F172A')
    )
    th_style = ParagraphStyle(
        f'ThStyle_{uid}',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=colors.white,
        alignment=1
    )
    td_style = ParagraphStyle(
        f'TdStyle_{uid}',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        textColor=colors.HexColor('#1E293B'),
        leading=9
    )
    td_center_style = ParagraphStyle(
        f'TdCenterStyle_{uid}',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        textColor=colors.HexColor('#1E293B'),
        alignment=1,
        leading=9
    )

    story = []
    story.append(Paragraph("BIMEJ 12 - BATALLÓN DE INGENIEROS DE MOVILIDAD Y CONTRAMOVILIDAD N° 12", title_style))

    periodo_label = mes if mes and mes != "TODOS" else "AÑO COMPLETO"
    if fecha_inicio or fecha_fin:
        periodo_label = f"{fecha_inicio or 'Inicio'} a {fecha_fin or 'Actual'}"

    if modo == "agil":
        story.append(Paragraph("REPORTE ÁGIL DE NOVEDADES (RESUMEN PARA COMANDANCIA - SIN DUPLICADOS)", subtitle_style))

        # Agrupar por integrante
        person_dict = defaultdict(list)
        for r in rows:
            person_key = (r[0], r[1], r[2])
            person_dict[person_key].append(r)

        items_list = []
        for (cedula, nombre, p_estado), p_rows in person_dict.items():
            dias_acumulados = len(p_rows)
            if min_dias > 0 and dias_acumulados < min_dias:
                continue

            p_rows.sort(key=lambda x: x[4])
            recs_for_ranges = [(x[4], x[3], x[7]) for x in p_rows]
            rango_html = format_agil_date_ranges(recs_for_ranges, highlight_html=True)

            subnovs = sorted(list(set(x[3] for x in p_rows)))
            subnov_str = ", ".join(subnovs)

            unique_descs = [x[7].strip() for x in p_rows if x[7] and str(x[7]).strip()]
            seen_descs = []
            for d in unique_descs:
                if d not in seen_descs:
                    seen_descs.append(d)
            desc_str = "; ".join(seen_descs[:2]) if seen_descs else "-"
            if len(desc_str) > 65:
                desc_str = desc_str[:62] + "..."

            items_list.append({
                "cedula": cedula,
                "nombre": nombre,
                "estado": p_estado,
                "sub_novedad": subnov_str,
                "rango_html": rango_html,
                "dias_acumulados": dias_acumulados,
                "descripcion": desc_str,
                "primera_fecha": p_rows[0][4],
                "ultima_fecha": p_rows[-1][4]
            })

        if orden == "nombre_asc":
            items_list.sort(key=lambda x: x["nombre"])
        elif orden == "dias_desc":
            items_list.sort(key=lambda x: x["dias_acumulados"], reverse=True)
        elif orden == "fecha_desc":
            items_list.sort(key=lambda x: x["ultima_fecha"], reverse=True)
        elif orden == "fecha_asc":
            items_list.sort(key=lambda x: x["primera_fecha"])

        meta_data = [
            [
                Paragraph("<b>Novedad:</b>", meta_key_style), Paragraph(subnovedad_nombre, meta_val_style),
                Paragraph("<b>Total Efectivos:</b>", meta_key_style), Paragraph(f"<b>{len(items_list)}</b>", meta_val_style),
            ],
            [
                Paragraph("<b>Periodo:</b>", meta_key_style), Paragraph(periodo_label, meta_val_style),
                Paragraph("<b>Total Días:</b>", meta_key_style), Paragraph(f"<b>{sum(it['dias_acumulados'] for it in items_list)}</b>", meta_val_style),
            ],
            [
                Paragraph("<b>Estado:</b>", meta_key_style), Paragraph(estado, meta_val_style),
                Paragraph("<b>Generado:</b>", meta_key_style), Paragraph(datetime.now().strftime("%d/%m/%Y %H:%M"), meta_val_style),
            ]
        ]
        meta_table = Table(meta_data, colWidths=[80, 280, 90, 270])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(meta_table)
        story.append(Spacer(1, 10))

        table_data = [[
            Paragraph("CÉDULA", th_style),
            Paragraph("APELLIDOS Y NOMBRES", th_style),
            Paragraph("ESTADO", th_style),
            Paragraph("RANGOS CONDENSADOS DE NOVEDAD", th_style),
            Paragraph("DÍAS", th_style),
            Paragraph("OBSERVACIÓN / MOTIVO", th_style)
        ]]

        for it in items_list:
            table_data.append([
                Paragraph(str(it["cedula"]), td_center_style),
                Paragraph(it["nombre"], td_style),
                Paragraph(it["estado"], td_center_style),
                Paragraph(it["rango_html"], td_style),
                Paragraph(f"<b>{it['dias_acumulados']}</b>", td_center_style),
                Paragraph(it["descripcion"], td_style)
            ])

        col_widths = [65, 175, 55, 240, 45, 142]
        t = Table(table_data, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 4),
            ('TOPPADDING', (0, 0), (-1, 0), 4),
            ('BOTTOMPADDING', (0, 1), (-1, -1), 3),
            ('TOPPADDING', (0, 1), (-1, -1), 3),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ]))
        story.append(t)
        filename = f"Reporte_Agil_Novedades_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"

    else:
        # Modo detallado
        story.append(Paragraph("REPORTE DETALLADO DE NOVEDADES (AUDITORÍA LÍNEA A LÍNEA)", subtitle_style))

        if min_dias > 0:
            person_counts = defaultdict(int)
            for r in rows:
                person_counts[r[0]] += 1
            rows = [r for r in rows if person_counts[r[0]] >= min_dias]

        if orden == "nombre_asc":
            rows.sort(key=lambda x: (x[1], x[4]))
        elif orden == "dias_desc":
            person_counts = defaultdict(int)
            for r in rows:
                person_counts[r[0]] += 1
            rows.sort(key=lambda x: (person_counts[x[0]], x[1], x[4]), reverse=True)
        elif orden == "fecha_desc":
            rows.sort(key=lambda x: (x[4], x[1]), reverse=True)
        elif orden == "fecha_asc":
            rows.sort(key=lambda x: (x[4], x[1]))

        meta_data = [
            [
                Paragraph("<b>Novedad:</b>", meta_key_style), Paragraph(subnovedad_nombre, meta_val_style),
                Paragraph("<b>Total Registros:</b>", meta_key_style), Paragraph(f"<b>{len(rows)}</b>", meta_val_style),
            ],
            [
                Paragraph("<b>Periodo:</b>", meta_key_style), Paragraph(periodo_label, meta_val_style),
                Paragraph("<b>Personal Único:</b>", meta_key_style), Paragraph(f"<b>{len(set(r[0] for r in rows))}</b>", meta_val_style),
            ],
            [
                Paragraph("<b>Estado:</b>", meta_key_style), Paragraph(estado, meta_val_style),
                Paragraph("<b>Generado:</b>", meta_key_style), Paragraph(datetime.now().strftime("%d/%m/%Y %H:%M"), meta_val_style),
            ]
        ]
        meta_table = Table(meta_data, colWidths=[80, 280, 90, 270])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(meta_table)
        story.append(Spacer(1, 10))

        table_data = [[
            Paragraph("CÉDULA", th_style),
            Paragraph("APELLIDOS Y NOMBRES", th_style),
            Paragraph("ESTADO", th_style),
            Paragraph("NOVEDAD", th_style),
            Paragraph("FECHA REP.", th_style),
            Paragraph("DESDE", th_style),
            Paragraph("HASTA", th_style),
            Paragraph("OBSERVACIÓN", th_style),
        ]]

        for r in rows:
            cedula_str = str(r[0])
            nombre_str = str(r[1])
            estado_str = str(r[2])
            novedad_str = str(r[3])
            fecha_rep_str = str(r[4])
            desde_str = str(r[5]) if r[5] else "-"
            hasta_str = str(r[6]) if r[6] else "-"
            obs_str = str(r[7]) if r[7] else "-"
            if len(obs_str) > 55:
                obs_str = obs_str[:52] + "..."

            table_data.append([
                Paragraph(cedula_str, td_center_style),
                Paragraph(nombre_str, td_style),
                Paragraph(estado_str, td_center_style),
                Paragraph(novedad_str, td_style),
                Paragraph(fecha_rep_str, td_center_style),
                Paragraph(desde_str, td_center_style),
                Paragraph(hasta_str, td_center_style),
                Paragraph(obs_str, td_style)
            ])

        col_widths = [65, 160, 50, 110, 65, 55, 55, 162]
        t = Table(table_data, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 4),
            ('TOPPADDING', (0, 0), (-1, 0), 4),
            ('BOTTOMPADDING', (0, 1), (-1, -1), 3),
            ('TOPPADDING', (0, 1), (-1, -1), 3),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ]))
        story.append(t)
        filename = f"Reporte_Detallado_Novedades_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"

    doc.build(story)
    pdf_buffer.seek(0)

    return StreamingResponse(
        pdf_buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
